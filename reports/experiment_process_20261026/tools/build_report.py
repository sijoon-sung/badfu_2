"""Build the process report from reviewed prose and the preserved legacy records."""
import hashlib
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

sys.stdout.reconfigure(encoding='utf-8')
OUT=Path(__file__).resolve().parents[1]
ROOT=OUT.parents[1]
ORDER=['R01','D01','D02','D03','D04','D05','D06','D07','D08','D09','D10','D11','T01','D12','D13','D14','D15','D16','D17','P01','P02','P03','D18','D19','R02','N01']
LINK=re.compile(r'\[([^\]]+)\]\(([^)]+)\)')

def clean(s):
    return re.sub(r'\*\*|`','',s).replace('正常','정상')

def legacy_appendix():
    records=json.loads((OUT/'legacy_source_records.json').read_text(encoding='utf-8'))
    lines=['# 과거 세부 실험 기록','',
           '기존 PR 보고서의 탐지 37개, 정화 19개, 기전 9개, 논문 확인 1개와 보류 2개를 모두 연결했다. 아래 PR 번호는 본문의 연구 카드 번호와 다르다. 같은 모델과 요청을 여러 방법으로 검사한 기록이며 독립 학습 68회를 뜻하지 않는다. 당시 수치와 후속 정정은 함께 읽는다.','',
           '탐지 PR D01~D25의 다수 검사는 축소 BadFU형 기존 시드71~74와 같은 요청을 재사용했다. 유지 R과 삭제 F로 만든 두 진단 분기는 실제 삭제 후 모델이 아니다. FC 1스텝, Conv 100스텝, 정상 KL 기준 자료를 쓰는 조건도 구분했다. 단순 변화량 후보를 현재 주력으로 다시 추천하지 않는다.','',
           'PR D26은 2026년 10월 19일 16시 58분의 70요청 보고서다. PR D27~D34도 당시 조회 사본이다. 이후 감사에서 같은 84요청과 정상 144요청을 별도로 비교했다. 서로 다른 조회의 건수를 합치지 않는다.','']
    names={'detection':'탐지','purification':'정화','mechanism':'기전','literature':'논문 확인','pending':'보류'}
    for r in records:
        name=clean(r['title'])
        if r['old_id']=='D26' and r['group']=='detection': name='고정 패치 탐지의 non IID 규모 확장'
        lines += [f"## 과거 PR {r['old_id']} {name}",'',f"확인하려던 것: {clean(r['hypothesis'])}",'',f"실제로 한 일: {clean(r['method'])}",'','관측 결과','',clean(r['result']),'',f"[원 보고서](../../{r['source']})",'']
    (OUT/'LEGACY_EXPERIMENTS.md').write_text('\n'.join(lines),encoding='utf-8')
    return records

def coverage_appendix():
    inv=json.loads((OUT/'folder_coverage.json').read_text(encoding='utf-8'))
    lines=['# 전체 폴더와 보고서의 대응','',
           '97개는 연구 보관 경로의 직접 하위 폴더 수다. 그림과 증거 사본, 공통 자원과 문헌 폴더도 포함돼 있어 실험 횟수가 아니다. 기존 인덱스의 카드 연결 여부는 그대로 유지했다. 내용 판정 미재검토 표시는 이 문서 작성만으로 새로운 성능 검증을 했다고 바꾸지 않았다는 뜻이다.','',
           '| 원래 폴더 | 기존 카드 또는 확인 상태 |','|---|---|']
    for x in inv:
        state=', '.join(x['related_cards']) if x['related_cards'] else '내용 판정 미재검토'
        lines.append(f"| [{x['path']}](../../{x['path']}) | {state} |")
    (OUT/'FOLDER_COVERAGE.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')

def add_link(p,label,target):
    rid=p.part.relate_to(target,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rid)
    run=OxmlElement('w:r');props=OxmlElement('w:rPr')
    col=OxmlElement('w:color');col.set(qn('w:val'),'333333');props.append(col)
    run.append(props);t=OxmlElement('w:t');t.text=label;run.append(t);link.append(run);p._p.append(link)

def inline(p,text,source):
    pos=0
    for m in LINK.finditer(text):
        if m.start()>pos:p.add_run(clean(text[pos:m.start()]))
        target=m.group(2)
        if not re.match(r'^(https?|file):',target):target=(source.parent/target).resolve().as_uri()
        add_link(p,clean(m.group(1)),target);pos=m.end()
    if pos<len(text):p.add_run(clean(text[pos:]))

def set_table(table,rows):
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    cols=len(rows[0]);width=174/cols
    if cols>=3:
        widths=[60]+[(174-60)/(cols-1)]*(cols-1)
    elif cols==2:widths=[112,62]
    else:widths=[174]
    for i,row in enumerate(table.rows):
        trPr=row._tr.get_or_add_trPr()
        if i==0:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        nosplit=OxmlElement('w:cantSplit');trPr.append(nosplit)
        for j,cell in enumerate(row.cells):
            cell.width=Mm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E7E7' if i==0 else 'FFFFFF');tcPr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcPr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'65' if side in ['top','bottom'] else '85');el.set(qn('w:type'),'dxa');margins.append(el)
            tcPr.append(margins)
            cell.text=clean(rows[i][j])
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.07
                if i==0:p.paragraph_format.keep_with_next=True
                if j>0 and len(cell.text)<36:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.font.size=Pt(9.2);run.bold=i==0

def render_markdown(doc,path,page_start=False):
    lines=path.read_text(encoding='utf-8').splitlines();i=0
    while i<len(lines):
        s=lines[i].strip()
        if not s:i+=1;continue
        if s.startswith('|') and i+1<len(lines) and re.match(r'^\|[\s:|\-]+\|$',lines[i+1].strip()):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                ss=lines[i].strip()
                if not re.match(r'^\|[\s:|\-]+\|$',ss):rows.append([c.strip() for c in ss.strip('|').split('|')])
                i+=1
            assert all(len(r)==len(rows[0]) for r in rows),(path,rows)
            t=doc.add_table(rows=len(rows),cols=len(rows[0]));set_table(t,rows)
            doc.add_paragraph().paragraph_format.space_after=Pt(2)
            continue
        if s.startswith('# '):
            title=clean(s[2:]);p=doc.add_paragraph(title,style='Title' if path.name=='INTRODUCTION.md' else 'Heading 1')
            if page_start:p.paragraph_format.page_break_before=True
        elif s=='## 원 기록':
            source_links=[]
            i+=1
            while i<len(lines):
                entry=lines[i].strip()
                if entry.startswith('- '):source_links.append(entry[2:])
                elif entry.startswith('#'):break
                i+=1
            p=doc.add_paragraph(style='Source');inline(p,'원 기록  '+' · '.join(source_links),path)
            continue
        elif s.startswith('## '):
            doc.add_paragraph(clean(s[3:]),style='Heading 2')
        elif re.match(r'^\d+\. ',s):
            p=doc.add_paragraph(style='Step');inline(p,s,path)
        elif s=='관측 결과':
            p=doc.add_paragraph();p.add_run(s).bold=True;p.paragraph_format.keep_with_next=True
        elif s.startswith('[원 보고서]'):
            if doc.paragraphs:doc.paragraphs[-1].paragraph_format.keep_with_next=True
            p=doc.add_paragraph(style='Source');inline(p,s,path)
        elif s.startswith('- '):
            p=doc.add_paragraph(style='Source' if '원 기록' in '\n'.join(lines[max(0,i-20):i]) else 'Step');inline(p,s[2:],path)
        else:
            p=doc.add_paragraph();inline(p,s,path)
        i+=1

def font_style(style,name,size):
    style.font.name=name;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    rf=style.element.get_or_add_rPr();fonts=rf.find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');rf.insert(0,fonts)
    for script in ['ascii','hAnsi','eastAsia']:fonts.set(qn('w:'+script),name)

records=legacy_appendix();coverage_appendix()
doc=Document();sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=Mm(17);sec.bottom_margin=Mm(17);sec.left_margin=Mm(18);sec.right_margin=Mm(18)
for style,size in [('Normal',10.2),('Title',20),('Heading 1',15),('Heading 2',11.2)]:
    font_style(doc.styles[style],'맑은 고딕',size)
normal=doc.styles['Normal'].paragraph_format
normal.line_spacing=1.10;normal.space_after=Pt(4)
for sty in ['Title','Heading 1','Heading 2']:
    doc.styles[sty].paragraph_format.keep_with_next=True
    doc.styles[sty].paragraph_format.space_before=Pt(6 if sty=='Heading 2' else 14 if sty=='Heading 1' else 0)
    doc.styles[sty].paragraph_format.space_after=Pt(4)
for name,size in [('Step',10.2),('Source',8.6)]:
    style=doc.styles.add_style(name,1);style.base_style=doc.styles['Normal'];font_style(style,'맑은 고딕',size)
    style.paragraph_format.space_after=Pt(3 if name=='Step' else 2)
for style in doc.styles:
    if style.type==1:
        ppr=style.element.get_or_add_pPr()
        for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');ppr.append(snap)
title_ppr=doc.styles['Title'].element.get_or_add_pPr()
title_border=OxmlElement('w:pBdr')
for side in ['top','left','bottom','right','between','bar']:
    edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'nil');title_border.append(edge)
title_ppr.append(title_border)
grid=sec._sectPr.find(qn('w:docGrid'))
if grid is not None:grid.set(qn('w:type'),'default')
doc.core_properties.author='성시준';doc.core_properties.last_modified_by='성시준'
doc.core_properties.title='BadFU 실험 과정과 결과'
doc.core_properties.subject='실험의 시작 이유와 모델 조건 실행 순서 결과 및 판단'
doc.core_properties.comments=''
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

render_markdown(doc,OUT/'INTRODUCTION.md')
for card in ORDER:render_markdown(doc,OUT/'experiments'/f'{card}.md',card=='R01')
render_markdown(doc,OUT/'LEGACY_EXPERIMENTS.md',True)
render_markdown(doc,OUT/'ADDITIONAL_RECORDS.md',True)
for p in doc.paragraphs:
    if p.style.name=='Title':
        ppr=p._p.get_or_add_pPr()
        for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
        border=OxmlElement('w:pBdr');edge=OxmlElement('w:bottom');edge.set(qn('w:val'),'nil');border.append(edge);ppr.append(border)
# Folder paths are navigation metadata, kept in a separate readable source file.
path=OUT/'BadFU_실험과정_전체정리_20261026.docx';doc.save(path)

parts=[OUT/'INTRODUCTION.md']+[OUT/'experiments'/f'{x}.md' for x in ORDER]+[OUT/'LEGACY_EXPERIMENTS.md',OUT/'ADDITIONAL_RECORDS.md']
combined=[]
for p in parts:
    txt=p.read_text(encoding='utf-8')
    def rebase(m):
        target=m.group(2)
        if re.match(r'^https?:',target):return m.group(0)
        resolved=(p.parent/target).resolve()
        import os
        return f'[{m.group(1)}]({Path(os.path.relpath(resolved,OUT)).as_posix()})'
    combined.append(LINK.sub(rebase,txt))
(OUT/'FULL_REPORT.md').write_text('\n\n'.join(combined)+'\n',encoding='utf-8')
nav=['# BadFU 실험 과정 정리','',
     '[전체 Word 문서](BadFU_실험과정_전체정리_20261026.docx) · [한 문서로 읽기](FULL_REPORT.md) · [전체 폴더 대응](FOLDER_COVERAGE.md)','',
     '각 항목에 시작 이유, 사용 모델과 자료, 실행 순서, 결과와 판단을 적었다. 원 실험과 실패는 기존 경로에 보존했다. 26개 연구 카드와 과거 세부 기록 68개는 겹치는 실험을 포함하며 독립 실험 수로 더하지 않는다.','',
     '| 카드 | 실험 과정 |','|---|---|']
for card in ORDER:
    title=(OUT/'experiments'/f'{card}.md').read_text(encoding='utf-8').splitlines()[0][2:]
    nav.append(f'| {card} | [{title[len(card)+1:]}](experiments/{card}.md) |')
nav += ['', '[과거 세부 실험 68개](LEGACY_EXPERIMENTS.md) · [카드가 없던 추가 파일럿](ADDITIONAL_RECORDS.md)','',
        '97개 폴더 중 기존 카드가 연결된 49개와 내용 판정 미재검토 48개의 구분은 유지했다. 새로 원 모델을 추론하거나 실험을 실행한 문서가 아니다.','', '[두 브랜치의 용도](BRANCHES.md)','']
(OUT/'README.md').write_text('\n'.join(nav),encoding='utf-8')
manifest={'created_at_utc':datetime.now(timezone.utc).isoformat(),'cards':ORDER,'legacy_records':len(records),'docx':path.name,'prose_files':[{ 'path':p.relative_to(OUT).as_posix(),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in parts]}
(OUT/'BUILD_MANIFEST.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'docx':str(path),'cards':len(ORDER),'legacy_records':len(records),'characters':sum(len(p.read_text(encoding='utf-8')) for p in parts)},ensure_ascii=False))
